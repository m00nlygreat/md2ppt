from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from pptx import Presentation


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class ExistingPptxCliTests(unittest.TestCase):
    def test_serve_help_exposes_theme_option(self):
        completed = subprocess.run(
            [sys.executable, str(PROJECT_ROOT / "main.py"), "serve", "--help"],
            cwd=PROJECT_ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=10,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertIn("--theme", completed.stdout)

    def test_template_ref_no_toc_and_debug_options_still_work(self):
        with tempfile.TemporaryDirectory(prefix="md2ppt-cli-") as directory:
            work = Path(directory)
            source = work / "간단한 발표.md"
            source.write_text(
                "---\ntitle: 회귀 확인\nauthor: 테스트\n---\n# 첫 장\n본문입니다.\n",
                encoding="utf-8",
            )
            debug_dir = work / "debug"
            commands = [
                [
                    sys.executable, str(PROJECT_ROOT / "main.py"), "-i", str(source),
                    "-o", str(work / "default.pptx"), "--no-toc", "--debug",
                    "--debug-dir", str(debug_dir),
                ],
                [
                    sys.executable, str(PROJECT_ROOT / "main.py"), "-i", str(source),
                    "-o", str(work / "named-template.pptx"), "--template", "hanaro",
                ],
                [
                    sys.executable, str(PROJECT_ROOT / "main.py"), "-i", str(source),
                    "-o", str(work / "reference.pptx"), "--ref",
                    str(PROJECT_ROOT / "md2ppt" / "refs" / "default.pptx"),
                ],
            ]
            for command in commands:
                completed = subprocess.run(
                    command,
                    cwd=PROJECT_ROOT,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=30,
                )
                self.assertEqual(
                    completed.returncode,
                    0,
                    (completed.stderr or "") + (completed.stdout or ""),
                )

            for output in ("default.pptx", "named-template.pptx", "reference.pptx"):
                output_path = work / output
                self.assertTrue(output_path.is_file())
                with zipfile.ZipFile(output_path) as package:
                    self.assertIn("ppt/presentation.xml", package.namelist())
            self.assertEqual(len(Presentation(work / "default.pptx").slides), 2)
            self.assertEqual(len(Presentation(work / "named-template.pptx").slides), 3)
            self.assertEqual(len(Presentation(work / "reference.pptx").slides), 3)
            for debug_file in (
                "0_flattened.md",
                "1_markdown_to_json.json",
                "2_json_to_slide.json",
            ):
                self.assertTrue((debug_dir / debug_file).is_file())
            self.assertEqual(
                (debug_dir / "0_flattened.md").read_text(encoding="utf-8"),
                source.read_text(encoding="utf-8").rstrip("\n"),
            )

            listed = subprocess.run(
                [sys.executable, str(PROJECT_ROOT / "main.py"), "--list-templates"],
                cwd=PROJECT_ROOT,
                capture_output=True,
                text=True,
                encoding="utf-8",
                timeout=10,
            )
            self.assertEqual(listed.returncode, 0, listed.stderr)
            self.assertIn("default", listed.stdout.splitlines())
            self.assertIn("hanaro", listed.stdout.splitlines())


if __name__ == "__main__":
    unittest.main()
